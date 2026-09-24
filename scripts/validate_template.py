#!/usr/bin/env python3
"""Validate that the bundled ledger is blank, portable, and structurally ready."""

from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile, ZipFile

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required to validate the XLSX template") from exc


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "assets" / "finance-os-template.xlsx"
REQUIRED_SHEETS = [
    "00 Dashboard",
    "01 Rules",
    "02 Balance Sheet",
    "03 Monthly History",
    "04 Investment Flow",
]
SENSITIVE_MARKERS = (
    "VO" + "O",
    "VB" + "IL",
    "Horizon " + "Robotics",
    "u" + "SMART",
    "161" + "725",
    "mango-" + "robotics",
    "docs.google." + "com",
)


def nonblank(sheet: object, cell_range: str) -> list[str]:
    values: list[str] = []
    for row in sheet[cell_range]:
        for cell in row:
            if cell.value not in (None, ""):
                values.append(str(cell.value))
    return values


def main() -> None:
    formulas = load_workbook(TEMPLATE, data_only=False, read_only=False)
    values = load_workbook(TEMPLATE, data_only=True, read_only=False)
    issues: list[str] = []

    if formulas.sheetnames[:5] != REQUIRED_SHEETS:
        issues.append("template sheets are missing or out of canonical order")

    rules_f = formulas["01 Rules"]
    if rules_f["A8"].value != "Monthly [TICKER] Contribution ([CCY])":
        issues.append("template passive-policy placeholder is missing")
    if any(rules_f.cell(row, 2).value not in (None, "") for row in range(5, 9)):
        issues.append("template contains a personal constraint value")
    if nonblank(values["01 Rules"], "A12:I67"):
        issues.append("template contains an active-exception record")

    if values["02 Balance Sheet"]["B2"].value not in (None, ""):
        issues.append("template contains a snapshot date")
    if nonblank(values["02 Balance Sheet"], "A5:G398") or nonblank(
        values["02 Balance Sheet"], "I5:I398"
    ):
        issues.append("template contains a balance-sheet input")
    if not str(formulas["02 Balance Sheet"]["H5"].value or "").startswith("="):
        issues.append("template balance-sheet value formula is missing")

    if nonblank(values["03 Monthly History"], "A5:E404") or nonblank(
        values["03 Monthly History"], "K5:K404"
    ):
        issues.append("template contains a monthly-history input")
    for coordinate in ("F5", "G5", "H5", "I5", "J5"):
        if not str(formulas["03 Monthly History"][coordinate].value or "").startswith("="):
            issues.append(f"template monthly formula is missing at {coordinate}")

    if nonblank(values["04 Investment Flow"], "A5:I406"):
        issues.append("template contains an investment event")

    for sheet in values.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(
                    marker.lower() in cell.value.lower() for marker in SENSITIVE_MARKERS
                ):
                    issues.append(f"template contains a prohibited marker at {sheet.title}!{cell.coordinate}")

    try:
        with ZipFile(TEMPLATE) as archive:
            for member in archive.namelist():
                normalized = member.lower()
                if normalized.startswith("xl/externallinks/"):
                    issues.append(f"template contains an external-link part: {member}")
                if not normalized.endswith((".xml", ".rels")):
                    continue
                payload = archive.read(member).decode("utf-8", errors="ignore").lower()
                for marker in SENSITIVE_MARKERS:
                    if marker.lower() in payload:
                        issues.append(f"template package contains a prohibited marker in {member}")
    except (BadZipFile, OSError) as exc:
        issues.append(f"template package cannot be inspected: {exc}")

    formulas.close()
    values.close()
    if issues:
        for issue in issues:
            print(f"FAIL: {issue}")
        raise SystemExit(1)
    print("blank ledger template: PASS")


if __name__ == "__main__":
    main()
